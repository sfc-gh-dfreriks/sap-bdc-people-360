#!/usr/bin/env python3
"""Build the SAP People 360 presales kit documents.

Mirrors the Finance 360 and Sales 360 builders so the five kits read as a set:

    00_START_HERE.docx              what is in the kit and which file to open
    03_SE_Quick_Start.docx          positioning, demo path, objections
    05_Architecture_and_Install.docx the medallion stack and how to stand it up
    06_Setup_and_Access.docx        the three regions, the listing, access

Every figure comes from /tmp/people_facts.json, produced by tools/people_facts.py
against the live account. Nothing is transcribed by hand.

    python3 tools/people_facts.py && python3 tools/build_presales_kit.py
"""
from __future__ import annotations

import json
import pathlib
import sys
from datetime import date

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from docx_kit import (  # noqa: E402
    AMBER,
    GREEN,
    GREY,
    RED,
    SAP_NAVY,
    SNOW_BLUE,
    body,
    bullet,
    callout,
    h1,
    h2,
    setup_page,
    table,
)

KIT = pathlib.Path.home() / "Documents" / "SAP" / "People_360_Presales_Kit"
FACTS_FILE = pathlib.Path("/tmp/people_facts.json")
DATE = date.today().strftime("%d %B %Y")

PAGE_PURPOSE = {
    "Workforce Overview": "Active headcount, attrition, average salary, gender split",
    "Headcount": "Headcount by department, company, tenure and age band",
    "Diversity & Inclusion": "Gender and generation mix, representation by department",
    "Compensation": "Salary distribution, compa-ratio against range midpoint",
    "Attrition": "Terminations by year, reason and department",
    "Performance": "Review rating distribution and coverage",
    "Learning": "Course completions, scores and status",
    "Recruiting": "Requisitions by status, open versus filled",
    "Org & Span": "Manager span of control and reporting depth",
    "Employees": "The record-level table, filterable",
    "BDC Sources & Lineage": "Which SAP BDC objects feed each layer",
    "Ask the Agent": "Natural-language questions over the workforce semantic view",
}


def load_facts() -> dict:
    if not FACTS_FILE.exists():
        sys.exit(f"{FACTS_FILE} missing — run tools/people_facts.py first")
    return json.loads(FACTS_FILE.read_text())["facts"]


def title_block(doc, title, subtitle, strap):
    for text_, size, bold, color, after in (
        (title, 20, True, SAP_NAVY, 2),
        (subtitle, 11.5, False, SNOW_BLUE, 2),
        (strap, 9, False, GREY, 14),
    ):
        p = doc.add_paragraph()
        r = p.add_run(text_)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)


def window_line(F) -> str:
    w = F["data_windows"]
    return (f"hires {w['DT_WORKFORCE_360']['min']} to {w['DT_WORKFORCE_360']['max']}, "
            f"learning to {w['DT_LEARNING']['max']}, "
            f"recruiting to {w['DT_RECRUITING']['max']}, "
            f"performance reviews {w['DT_PERFORMANCE']['min']}–{w['DT_PERFORMANCE']['max']}")


def pii_callout(doc, F):
    ids = ", ".join(r["EMPLOYEE_ID"] for r in F["sample_employee_ids"])
    callout(
        doc,
        "Ask this before anyone else does",
        f"This is workforce data, so the first question will be whether these are "
        f"real people. They are not. The source is a demo database "
        f"(SAP_BDC_DEMO_CORE_WORKFORCE_DATA), employee identifiers are synthetic "
        f"sequential IDs ({ids}…) rather than SAP personnel numbers, and no names "
        f"are surfaced anywhere in the analytics layer. Say so up front — it "
        f"settles the room and it is true.",
    )


def agent_scope_callout(doc, F):
    sv = next(iter(F["semantic_view_detail"].values()))
    callout(
        doc,
        "The one limitation to know",
        f"Ask the Agent answers on workforce only. The semantic view carries a "
        f"single table ({', '.join(sv['table_names'])}) with {sv['dimensions']} "
        f"dimensions and {sv['facts']} facts — so headcount, attrition, "
        f"compensation, diversity and tenure all work, but performance, learning "
        f"and recruiting are dashboard-only. All eight built-in questions are "
        f"workforce questions, which is the app telling you the same thing.",
    )


# ---------------------------------------------------------------- START HERE


def build_start_here(F):
    doc = Document()
    setup_page(doc)
    title_block(
        doc,
        "SAP People 360",
        "Presales kit — start here",
        f"SAP Partnership Compass  ·  {DATE}  ·  Owner: Dave Freriks",
    )

    body(
        doc,
        "People 360 shows SAP SuccessFactors workforce data working as analytics on "
        "Snowflake, reached through SAP Business Data Cloud with zero copy. Headcount, "
        "attrition, compensation, diversity, performance, learning and recruiting in "
        "one governed model, with a Cortex Agent over the workforce layer. It installs "
        "as a self-contained Native App — nothing to configure, no data to load.",
    )

    callout(
        doc,
        "Fastest path to a demo",
        f"Install the Native App listing {F['app_listing']} from the internal "
        f"Marketplace in your nearest region, open it, and you have a "
        f"{F['app_page_count']}-page workforce dashboard with its own data already "
        f"inside. See 06_Setup_and_Access.",
    )

    pii_callout(doc, F)

    k = F["kpis"]
    sv = next(iter(F["semantic_view_detail"].values()))
    table(
        doc,
        ["The demo in eight numbers", "Value"],
        [
            ["Employee records", f"{F['employees']:,} ({k['active_headcount']:,} active, {F['terminations']} terminated)"],
            ["Attrition shown", f"{k['attrition_pct']}% — terminations as a share of all records, not annualised"],
            ["Organisation", f"{F['departments']} departments across {F['companies']} companies"],
            ["Compensation", f"${k['avg_salary']:,} average base, compa-ratio {F['avg_compa_ratio']}"],
            ["Workforce mix", f"{k['pct_female']}% female, {k['pct_managers']}% managers, {F['critical_positions']} critical positions"],
            ["Analytics layer", f"{len(F['analytics_objects'])} tables, {F['analytics_rows']:,} rows"],
            ["Semantic view", f"{sv['tables']} table, {sv['dimensions']} dimensions, {sv['facts']} facts"],
            ["App", f"{F['app_page_count']} pages, published in {len(F['regions'])} regions"],
        ],
        widths=[2.2, 4.5],
    )
    body(
        doc,
        f"Every figure above was read from the account on {F['verified_on']}, not "
        f"copied from a previous deck. The KPI row quotes the application's own "
        f"/api/overview so the kit cannot disagree with the screen.",
        italic=True,
    )

    h1(doc, "Which file to open")
    table(
        doc,
        ["File", "Use it when", "Read time"],
        [
            ["03_SE_Quick_Start.docx",
             "You are demoing this week. Positioning, a ten-minute path, the eight "
             "agent questions, discovery questions, objections. Start here.", "10 min"],
            ["02_Demo_Scripts_by_Persona.docx",
             "You know who is in the room — CHRO, HR ops, talent, finance, architect. "
             "Each script opens on a different page.", "per script"],
            ["01_Management_Summary.docx",
             "Leaving something with the customer, or briefing an exec beforehand. "
             "Customer-safe and explicit about what is real.", "15 min"],
            ["00_Presales_Overview.pptx",
             "You need slides. Ten slides: the problem, the architecture, the app, "
             "the agent and its limits, regions, and the ask.", "10 slides"],
            ["05_Architecture_and_Install.docx",
             "An architect is in the room, or you are standing it up yourself.", "reference"],
            ["06_Setup_and_Access.docx",
             "You need access: the listing, the three regions, the public build.", "5 min"],
        ],
        widths=[2.0, 3.9, 0.8],
    )

    agent_scope_callout(doc, F)

    h1(doc, "One thing that is different about HR demos")
    body(
        doc,
        "Workforce analytics lands differently from supply chain or finance. Pay, "
        "performance ratings and diversity are sensitive even when synthetic, and in "
        "EMEA a works council may have a say in what an HR analytics platform is "
        "allowed to show. Two habits help: say the data is synthetic before you are "
        "asked, and talk about governance — row access policies, masking, who can see "
        "compensation — as part of the demo rather than as an afterthought. That is a "
        "Snowflake strength, so it is worth raising rather than avoiding.",
    )

    h1(doc, "The companion kits")
    body(
        doc,
        "Four sibling kits follow the same pattern on the same Compass page: Supply "
        "Chain 360, Supply Chain Ontology, Finance 360 and Sales 360. Same "
        "architecture, same BDC zero-copy story, different domain. If a customer is "
        "running SuccessFactors alongside S/4HANA, People 360 plus one other kit "
        "makes the cross-domain point better than either alone.",
    )

    h1(doc, "Support")
    body(doc, f"Source and documentation: {F['repo']}. Questions: Dave Freriks.")

    out = KIT / "00_START_HERE.docx"
    doc.save(out)
    return out


# --------------------------------------------------------------- QUICK START


def build_quick_start(F):
    doc = Document()
    setup_page(doc)
    title_block(
        doc,
        "SAP People 360",
        "SE quick start",
        f"Ten-minute demo path, the agent's scope, and objection handling  ·  {DATE}",
    )

    body(doc, "Read once, the day before you demo. It assumes you have not opened the app.")

    h1(doc, "Positioning, in three sentences")
    bullet(
        doc,
        "HR leaders are asked questions their systems answer slowly — where attrition "
        "is concentrated, whether pay is equitable, which critical roles are exposed — "
        "and the data sits in SuccessFactors behind an extract request.",
    )
    bullet(
        doc,
        "SAP Business Data Cloud shares that workforce data into Snowflake with zero "
        "copy, keeping SAP's own definitions of headcount, tenure and compa-ratio "
        "rather than having an analyst rebuild them in a spreadsheet.",
    )
    bullet(
        doc,
        "Snowflake adds the governed semantic layer, the natural-language agent and "
        "the app — and the governance controls that make HR data safe to expose at all.",
    )

    h1(doc, "Before you start")
    bullet(doc, f"Install the Native App ({F['app_listing']}) in your nearest region, or open the public build to rehearse. Open it once to warm the container.")
    bullet(doc, "Check the Department and Company filters — they apply across every page, and a stray selection makes charts look empty.")
    bullet(doc, f"Know the windows: {window_line(F)}. Performance is a year behind the rest, so do not call it current.")
    bullet(doc, "Decide how you will handle the synthetic-data question. Raise it yourself in the first minute; do not wait to be challenged on HR data.")

    h1(doc, "The ten-minute path")
    k = F["kpis"]
    table(
        doc,
        ["#", "Page", "Do and say", "Time"],
        [
            ["1", "—",
             "Open the app cold. Say: nothing installed, nothing loaded, and this is "
             "synthetic workforce data modelled on real SuccessFactors structure.", "1 min"],
            ["2", "Workforce Overview",
             f"{k['active_headcount']:,} active, {k['attrition_pct']}% attrition, "
             f"${k['avg_salary']:,} average base, {k['pct_female']}% female. Say these "
             f"are SAP's definitions, not ones we invented downstream.", "1 min"],
            ["3", "Attrition",
             "Terminations by year, reason and department. This is the page HR "
             "leaders lean into — where is it concentrated, and why.", "2 min"],
            ["4", "Compensation",
             f"Compa-ratio against range midpoint, average {F['avg_compa_ratio']}. "
             f"Pay equity is the question behind the question here.", "1 min"],
            ["5", "Diversity & Inclusion",
             "Gender and generation mix by department. Handle plainly and without "
             "editorialising; the value is that it is measurable at all.", "1 min"],
            ["6", "Org & Span",
             f"Span of control, with {F['critical_positions']} critical positions "
             f"flagged. Good bridge to succession-risk conversations.", "1 min"],
            ["7", "Ask the Agent",
             "Ask one of the eight built-in questions live, then expand the generated "
             "SQL — governed and explainable, not a black box.", "2 min"],
            ["8", "BDC Sources & Lineage",
             "Close here for a technical room: it shows which SAP BDC objects feed "
             "each layer, which is the zero-copy proof.", "1 min"],
        ],
        widths=[0.3, 1.5, 4.2, 0.7],
    )

    h2(doc, "The eight questions the agent ships with")
    body(doc, "Read from the app's source, so these are exactly the chips on screen.")
    for q in F["agent_questions"]:
        bullet(doc, q)
    agent_scope_callout(doc, F)

    h2(doc, "The closing line")
    callout(
        doc,
        "Say this",
        "Everything here reads SuccessFactors data in place — no extract, no copy, no "
        "re-derived HR metrics — and the governance that decides who may see "
        "compensation is enforced in the platform rather than in a spreadsheet's "
        "sharing settings.",
    )

    h1(doc, "What is on each page")
    table(
        doc,
        ["Page", "What it shows"],
        [[p, PAGE_PURPOSE.get(p, "—")] for p in F["app_pages"]],
        widths=[1.9, 4.8],
    )

    h1(doc, "Discovery questions")
    bullet(doc, "How long does it take today to answer where attrition is concentrated, and who does that work?")
    bullet(doc, "Is workforce reporting done in SuccessFactors, in spreadsheets, or somewhere else entirely?")
    bullet(doc, "Do you already have a pay-equity review cycle, and what does it run on?")
    bullet(doc, "Who is allowed to see compensation at record level, and how is that enforced now?")
    bullet(doc, "Are you on RISE with SAP, and is BDC already part of that subscription?")
    bullet(doc, "Is there a works council or data-protection review that any HR analytics change has to clear?")

    h1(doc, "Objections, and what to say")
    table(
        doc,
        ["They say", "You say"],
        [
            ["Is this real employee data?",
             "No. It is synthetic, in a demo database, with sequential IDs and no "
             "names in the analytics layer. The structure is modelled on real "
             "SuccessFactors fields; the people are not real."],
            ["Our HR data cannot leave SAP.",
             "It does not. BDC shares it zero-copy into Snowflake — there is no "
             "extract and no second copy to secure. That is usually the objection "
             "that BDC exists to answer."],
            ["Who would be able to see salaries?",
             "Whoever you decide. Masking policies and row access policies are "
             "enforced in Snowflake, so an HR business partner can see their own "
             "population and nothing else. Offer to show it rather than assert it."],
            ["Can the agent answer anything about performance?",
             "Not today. The semantic view covers the workforce table only, so the "
             "agent handles headcount, attrition, pay, diversity and tenure. "
             "Performance, learning and recruiting are on the dashboards. Extending "
             "the model is straightforward and is a good scoped follow-on."],
            ["Why do hires only appear in the last two years?",
             "An artefact of the synthetic dataset: the hire-count field is only "
             "populated for recent periods, so the chart shows terminations across "
             "the full history but hires mostly in 2025 and 2026. Say it is a "
             "dataset artefact rather than improvising an explanation."],
            ["How current is this?",
             f"It varies by table, which is worth being precise about: {window_line(F)}."],
            ["Can we point it at our own SuccessFactors data?",
             "Yes, and that is the follow-on. The medallion SQL, the semantic view "
             "and the agent are the deliverable; swap the L0 layer for the "
             "customer's own BDC workforce data products."],
        ],
        widths=[1.9, 4.8],
    )

    h1(doc, "If it goes wrong")
    bullet(doc, "A page is empty — check the Department and Company filters; they persist across pages.")
    bullet(doc, "The agent returns nothing useful — you probably asked about performance, learning or recruiting. It only knows workforce.")
    bullet(doc, "The app is slow on first open — the container is cold. Open it before the call.")
    bullet(doc, "A number disagrees with this document — trust the screen and tell me; the kit is generated from the account and should be regenerated.")

    out = KIT / "03_SE_Quick_Start.docx"
    doc.save(out)
    return out


# ------------------------------------------------- ARCHITECTURE AND INSTALL


def build_architecture(F):
    doc = Document()
    setup_page(doc)
    title_block(
        doc,
        "SAP People 360",
        "Architecture and install",
        f"The medallion stack, every object, and the build order  ·  {DATE}",
    )

    body(
        doc,
        "Everything sits inside Snowflake and reads SAP workforce data through SAP "
        "Business Data Cloud zero-copy shares. Object counts below were read from the "
        f"account on {F['verified_on']}.",
    )

    h1(doc, "The stack")
    l0 = ", ".join(F["app_reads_share_directly"]) or "SAP_BDC_DEMO_CORE_WORKFORCE_DATA"
    sv_fqn, sv = next(iter(F["semantic_view_detail"].items()))
    table(
        doc,
        ["Layer", "Objects", "What it does"],
        [
            ["L0 Bronze", "SAP_BDC_DEMO_CORE_WORKFORCE_DATA",
             f"The SAP BDC workforce data product, shared zero-copy. The app reads one "
             f"object here directly for the lineage page: {l0}."],
            ["L1 Silver", f"{len(F['l1_objects'])} view in SAP_BDC_L1",
             "SAP_BDC_L1.WORKFORCE — a typed, renamed projection over the share. "
             "Unlike Supply Chain 360, People 360 does keep an explicit L1."],
            ["L2 Gold", f"{len(F['analytics_objects'])} tables in ANALYTICS, {F['analytics_rows']:,} rows",
             "DT_WORKFORCE_360 plus DT_PERFORMANCE, DT_LEARNING and DT_RECRUITING."],
            ["Semantic", sv_fqn.split(".", 1)[1],
             f"{sv['tables']} table ({', '.join(sv['table_names'])}), "
             f"{sv['dimensions']} dimensions, {sv['facts']} facts, "
             f"{sv['verified_queries']} verified queries."],
            ["Agent", F["agents"][0].split(".", 1)[1] if F["agents"] else "—",
             "Cortex Agent for natural-language workforce questions, and the model "
             "behind the app's Ask the Agent page."],
            ["App", f"{F['app_page_count']} pages",
             "React client and Express server, packaged as a Native App."],
        ],
        widths=[0.9, 2.1, 3.7],
    )

    h1(doc, "Only one of the four gold tables is a dynamic table")
    body(
        doc,
        f"A DT_ prefix is not evidence. SHOW DYNAMIC TABLES returns exactly one: "
        f"{', '.join(F['dynamic_tables'])}. The other three — "
        f"{', '.join(F['dt_prefixed_not_dynamic'])} — are plain base tables that carry "
        f"the prefix by convention and never refresh on their own.",
    )
    dtl = F["dynamic_table_detail"][0] if F["dynamic_table_detail"] else {}
    if dtl:
        table(
            doc,
            ["Dynamic table", "Target lag", "Refresh mode", "Last refreshed"],
            [[dtl.get("name"), dtl.get("target_lag"), dtl.get("refresh_mode"),
              dtl.get("data_timestamp")]],
            widths=[2.2, 1.5, 1.5, 1.5],
        )
        body(
            doc,
            f"Target lag is {dtl.get('target_lag')}, which means it refreshes only when "
            f"something downstream with its own lag asks for it. It last refreshed "
            f"{dtl.get('data_timestamp')}. If you point this stack at live BDC data, set "
            f"an explicit TARGET_LAG and promote the other three to dynamic tables, or "
            f"they will silently stay frozen.",
        )

    h1(doc, "Data windows, per table")
    body(doc, "The four tables do not share a date column, so there is no single window. State them separately.")
    rows = []
    for name, w in F["data_windows"].items():
        span = (f"{w['min']} to {w['max']}" if w["kind"] == "date"
                else f"review years {w['min']}–{w['max']}")
        rows.append([name, w["column"], span, f"{w.get('non_null', '—')}"])
    table(doc, ["Table", "Column", "Span", "Non-null"], rows, widths=[2.4, 1.7, 1.8, 0.8])
    body(
        doc,
        "Performance reviews stop a year before the rest of the dataset. That is the "
        "one window most likely to catch someone out in a demo.",
        italic=True,
    )

    h1(doc, "What reads what")
    table(
        doc,
        ["Consumer", "Reads"],
        [
            ["The application", ", ".join(sorted(F["app_data_sources"])) or "—"],
            ["Cortex Analyst", F.get("analyst_semantic_view") or "—"],
        ],
        widths=[1.5, 5.2],
    )
    body(
        doc,
        "The app is a hybrid: most pages read the L2 tables, and the lineage page "
        "reads the raw BDC share directly to prove the zero-copy path. It does not "
        "read the semantic view — that is the agent's, not the dashboard's.",
    )

    h1(doc, "Build order")
    table(
        doc,
        ["Step", "What to run", "Result"],
        [
            ["1", "sql/ — L1 view over the BDC share", "SAP_BDC_L1.WORKFORCE"],
            ["2", "sql/ — L2 analytics", "DT_WORKFORCE_360 and the three base tables"],
            ["3", "sql/ — semantic view", sv_fqn.split(".", 1)[1]],
            ["4", "sql/ — agent", F["agents"][0].split(".", 1)[1] if F["agents"] else "—"],
            ["5", "scripts/ — build, migrate, deploy", "Native App package and version"],
            ["6", "scripts/ — org listing", "Region-scoped organization listing"],
        ],
        widths=[0.5, 2.9, 3.3],
    )
    body(doc, "After step 4 you can demo through Snowflake Intelligence without the app at all.")

    h1(doc, "A check worth doing before you rely on the listing")
    sv_ext = F["semantic_views"][0].get("extension") if F["semantic_views"] else None
    body(
        doc,
        f"The semantic view reports extension {sv_ext!r}. On the Supply Chain 360 and "
        f"Ontology views this field carries CA and/or AI, which is what lets a "
        f"marketplace listing attach the view and agent as AI products. Cortex Analyst "
        f"works here regardless because the app passes the view explicitly, but if you "
        f"intend to attach AI products to a People 360 data-share listing, verify that "
        f"first rather than assuming parity with the other domains.",
    )

    out = KIT / "05_Architecture_and_Install.docx"
    doc.save(out)
    return out


# ------------------------------------------------------------- SETUP, ACCESS


def build_setup(F):
    doc = Document()
    setup_page(doc)
    title_block(
        doc,
        "SAP People 360",
        "Setup and access",
        f"The listing, the three regions, and the public build  ·  {DATE}",
    )

    h1(doc, "How to get to it")
    table(
        doc,
        ["Route", "What you get", "Notes"],
        [
            ["Native App listing", f"The {F['app_page_count']}-page dashboard with data "
             "bundled in. No grants, no warehouse sizing, no data setup.",
             f"{F['app_listing']} — region-scoped, install the one for your region"],
            ["Public static build", "A credential-free browser build for rehearsal or "
             "async sharing. No Snowflake behind it, so Ask the Agent will not answer.",
             F["public_url"]],
            ["Your own account", "Run the SQL, then demo through Snowflake "
             "Intelligence with no app at all.", "See 05_Architecture_and_Install"],
        ],
        widths=[1.3, 3.0, 2.4],
    )
    body(
        doc,
        "Note that People 360 ships the Native App listing only. Unlike Supply Chain "
        "360, there is no separate data-share listing today, so there is no "
        "mount-the-tables-in-your-own-account route from the Marketplace.",
    )

    h1(doc, "Where it is published")
    rows = []
    for r in F["regions"]:
        if "error" in r:
            rows.append([r["connection"], "—", f"check failed: {r['error']}"])
            continue
        listings = ", ".join(f"{x['name']} ({x['state']})" for x in r["listings"]) or "none found"
        rows.append([r["connection"], r["region"], listings])
    table(doc, ["Connection", "Snowflake region", "Listing"], rows, widths=[1.7, 2.1, 2.9])
    body(
        doc,
        "Each region is an independent governed install of the same definition. There "
        "is no cross-region querying, and for workforce data that separation is "
        "usually a feature worth naming rather than a limitation to apologise for.",
        italic=True,
    )

    h1(doc, "Snowflake objects")
    sv_fqn, sv = next(iter(F["semantic_view_detail"].items()))
    table(
        doc,
        ["Object", "Detail"],
        [
            ["Database", "SAP_PEOPLE_360"],
            ["L0 share", "SAP_BDC_DEMO_CORE_WORKFORCE_DATA.BDCCONNECT"],
            ["L1", "SAP_BDC_L1.WORKFORCE (view)"],
            ["ANALYTICS", f"{len(F['analytics_objects'])} tables, {F['analytics_rows']:,} rows — "
                          f"{len(F['dynamic_tables'])} dynamic, "
                          f"{len(F['dt_prefixed_not_dynamic'])} plain base tables"],
            ["Semantic view", sv_fqn],
            ["Agent", F["agents"][0] if F["agents"] else "—"],
            ["Account checked", f"{F['account']} ({F['region']})"],
        ],
        widths=[1.6, 5.1],
    )
    callout(
        doc,
        "Agent location",
        f"The agent lives in the AGENTS schema — {F['agents'][0] if F['agents'] else '—'}. "
        f"That differs from Supply Chain 360, where the agent sits in ANALYTICS. If you "
        f"are adapting a script between domains, check the schema.",
    )

    h1(doc, "Privacy posture")
    pii_callout(doc, F)
    body(
        doc,
        "If a customer asks what production would look like, the honest answer is that "
        "it needs a governance design: masking policies on compensation, row access "
        "policies so managers see only their own population, and an agreed position "
        "with whoever reviews HR data changes. That conversation is a good outcome, "
        "not an obstacle — it is where Snowflake is strong.",
    )

    h1(doc, "Related assets")
    table(
        doc,
        ["Asset", "Where"],
        [
            ["Finance 360 kit", "Same pattern, finance domain — Compass"],
            ["Sales 360 kit", "Same pattern, sales domain — Compass"],
            ["Supply Chain 360 kit", "Same pattern, supply chain — Compass"],
            ["Supply Chain Ontology kit", "Disruption modelling companion — Compass"],
            ["This kit", "SAP Partnership Compass, Seismic"],
        ],
        widths=[2.0, 4.7],
    )

    h1(doc, "Contact")
    body(doc, "Dave Freriks — for access help, or to scope pointing this at a customer's own BDC workforce data.")

    out = KIT / "06_Setup_and_Access.docx"
    doc.save(out)
    return out


def main() -> int:
    F = load_facts()
    KIT.mkdir(parents=True, exist_ok=True)
    for fn in (build_start_here, build_quick_start, build_architecture, build_setup):
        print(f"wrote {fn(F)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
