#!/usr/bin/env python3
"""Build the SAP People 360 presales deck — ten slides on the SAP branded template.

Mirrors the Sales 360 and Finance 360 deck builders: same template, same helper
kit, same verifier. Figures come from /tmp/people_facts.json and screenshots from
/tmp/people_shots, so the deck cannot drift from the account or the application.

    python3 tools/people_facts.py
    python3 ../sap-bdc-finance-360/tools/capture_shots.py --app people
    python3 tools/build_presales_deck.py
"""

import json
import pathlib
import sys

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from pptx_kit import (  # noqa: E402
    BODY_GREY, DK1, DK2, LIGHT_BG, SF_BLUE, TEAL, VIOLET, WHITE,
    add_shape_text, new_presentation, set_ph, verify_deck, verify_slide,
)
from pptx.dml.color import RGBColor  # noqa: E402

RED = RGBColor(0xA2, 0x00, 0x00)          # in-palette dark red, safe as text

FACTS = pathlib.Path("/tmp/people_facts.json")
SHOTS = pathlib.Path("/tmp/people_shots")
OUT = (pathlib.Path.home() / "Documents" / "SAP"
       / "People_360_Presales_Kit" / "00_Presales_Overview.pptx")

TOP, BOTTOM, LEFT, RIGHT = 1.32, 5.08, 0.40, 9.50
FULLW = RIGHT - LEFT


# ------------------------------------------------------------------ facts


def load_facts():
    doc = json.loads(FACTS.read_text())
    f = doc["facts"]
    shots = {s["id"]: s["file"] for s in json.loads((SHOTS / "manifest.json").read_text())}
    for sid, path in shots.items():
        if not pathlib.Path(path).exists():
            raise FileNotFoundError(f"screenshot missing for '{sid}': {path}")
    return f, shots


# ----------------------------------------------------------------- helpers
# Copied from the Sales 360 builder so the two decks lay out identically.


def content(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    set_ph(s, 0, title)
    set_ph(s, 1, subtitle)
    return s


def box(slide, x, y, w, h, fill=LIGHT_BG, accent=None):
    add_shape_text(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, "", fill, DK1)
    if accent is not None:
        add_shape_text(slide, MSO_SHAPE.RECTANGLE, x, y, 0.05, h, "", accent, DK1)


def stack(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, spacing=1.06):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(runs):
        body, size, bold, colour = item[:4]
        after = item[4] if len(item) > 4 else 5
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        r = p.add_run()
        r.text = body
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Arial"
    return tb


def card(slide, x, y, w, h, kicker, runs, accent=DK2, fill=LIGHT_BG):
    box(slide, x, y, w, h, fill, accent)
    head = [(kicker.upper(), 8.5, True, DK2, 7)] if kicker else []
    stack(slide, x + 0.24, y + 0.16, w - 0.42, h - 0.30, head + list(runs))


def banner(slide, y, runs, fill=DK2, h=0.62, x=LEFT, w=FULLW):
    add_shape_text(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, "", fill, WHITE)
    stack(slide, x + 0.30, y + 0.11, w - 0.60, h - 0.22, runs, align=PP_ALIGN.LEFT)


def stat(slide, x, y, w, h, value, label, detail=None, fill=LIGHT_BG, accent=DK2):
    box(slide, x, y, w, h, fill, accent)
    runs = [(value, 28, True, DK2, 3), (label.upper(), 8.5, True, BODY_GREY, 4)]
    if detail:
        runs.append((detail, 10, False, DK1, 0))
    stack(slide, x + 0.26, y + 0.16, w - 0.46, h - 0.30, runs)


def note(slide, text, y=4.88):
    stack(slide, LEFT, y, FULLW, 0.19, [(text, 8, False, BODY_GREY, 0)])


def picture(slide, path, x, y, w, h):
    im = Image.open(path)
    ar = im.width / im.height
    if ar > w / h:
        pw, ph = w, w / ar
    else:
        ph, pw = h, h * ar
    slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
        Inches(pw), Inches(ph))


def caption(slide, x, y, w, text):
    stack(slide, x, y, w, 0.20, [(text.upper(), 8.5, True, DK2, 0)])


def region_label(raw: str) -> str:
    """AWS us-west-2 rather than PUBLIC.AWS_US_WEST_2.

    The raw identifier is a long all-caps token: the deck verifier flags it as
    merged caps, and it reads badly at 12pt. Cloud stays capitalised, the region
    goes lower-case with hyphens, which is how the provider writes it anyway.
    """
    name = raw.split(".")[-1]
    parts = name.split("_")
    if not parts:
        return raw
    cloud, rest = parts[0], parts[1:]
    return f"{cloud} {'-'.join(x.lower() for x in rest)}" if rest else cloud


def window_line(f):
    w = f["data_windows"]
    return (f"hires {w['DT_WORKFORCE_360']['min']}–{w['DT_WORKFORCE_360']['max']} · "
            f"learning to {w['DT_LEARNING']['max']} · "
            f"recruiting to {w['DT_RECRUITING']['max']} · "
            f"reviews {w['DT_PERFORMANCE']['min']}–{w['DT_PERFORMANCE']['max']}")


# ---------------------------------------------------------------------- slides


def s01_cover(prs, f):
    s = prs.slides.add_slide(prs.slide_layouts[13])
    set_ph(s, 3, "SAP PEOPLE 360")
    set_ph(s, 0, "SuccessFactors workforce data as governed analytics")
    set_ph(s, 2, f"SE presales kit · verified {f['verified_on']}")
    return s


def s02_problem(prs, f):
    s = content(prs, "The data exists. Reaching it is the project.",
                "Workforce questions that should take minutes take weeks")
    card(s, LEFT, TOP, 4.4, 2.5, "what HR is asked", [
        ("Where is attrition concentrated, and why?", 12, False, DK1, 8),
        ("Is our pay equitable across grade and gender?", 12, False, DK1, 8),
        ("Which critical roles are a single resignation away from a gap?", 12, False, DK1, 8),
        ("How does headcount actually break down by department and generation?", 12, False, DK1, 0),
    ], accent=SF_BLUE)
    card(s, LEFT + 4.7, TOP, 4.4, 2.5, "why it is slow", [
        ("The data sits in SuccessFactors, reachable by extract request.", 12, False, DK1, 8),
        ("Each extract gets remodelled in a spreadsheet, so headcount and tenure "
         "stop matching SAP's own definitions.", 12, False, DK1, 8),
        ("Nobody wants employee-level pay data copied into another system, so "
         "the analysis gets narrowed instead of governed.", 12, False, DK1, 0),
    ], accent=RED)
    banner(s, 4.05, [
        ("The blocker is not analytics capability. It is that moving HR data "
         "anywhere is the part nobody wants to own.", 12.5, True, WHITE, 0)])
    note(s, "SAP BDC removes the move: the data is shared into Snowflake zero-copy, "
            "with SAP's definitions intact.")
    return s


def s03_pattern(prs, f):
    sv_fqn, sv = next(iter(f["semantic_view_detail"].items()))
    s = content(prs, "One medallion stack, entirely inside Snowflake",
                "Zero copy in, governed model out")
    # Layer names are spaced for reading rather than written as raw FQNs: a long
    # ALL_CAPS_IDENTIFIER trips the deck verifier's merged-caps check and is hard
    # to read at 11pt. The exact fully-qualified names are in 06_Setup_and_Access.
    layers = [
        ("L0", "SAP BDC workforce data product", "Shared zero-copy, treated as immutable", TEAL),
        ("L1", "SAP BDC L1 · Workforce view", "Typed, renamed projection over the share", SF_BLUE),
        ("L2", f"Analytics — {len(f['analytics_objects'])} tables, {f['analytics_rows']:,} rows",
         f"{len(f['dynamic_tables'])} dynamic table, "
         f"{len(f['dt_prefixed_not_dynamic'])} plain base tables", SF_BLUE),
        ("SEM", "SAP People 360 Analytics semantic view",
         f"{sv['tables']} table · {sv['dimensions']} dimensions · {sv['facts']} facts", VIOLET),
        ("AI", "SAP People Analyst agent",
         "Cortex Agent, and Cortex Analyst inside the app", VIOLET),
        ("APP", f"Native App — {f['app_page_count']} pages",
         "React + Express, data bundled, nothing to configure", DK2),
    ]
    y = TOP
    for tag, name, detail, accent in layers:
        # TEAL is a light accent: white text on it fails the contrast check.
        tag_text = DK1 if accent is TEAL else WHITE
        add_shape_text(s, MSO_SHAPE.RECTANGLE, LEFT, y, 0.62, 0.52, tag, accent, tag_text, 11, True)
        box(s, LEFT + 0.68, y, FULLW - 0.68, 0.52, LIGHT_BG, None)
        stack(s, LEFT + 0.88, y + 0.07, FULLW - 1.1, 0.40, [
            (name, 11, True, DK1, 1), (detail, 9, False, BODY_GREY, 0)])
        y += 0.60
    note(s, "Nothing is extracted and no second copy is created. Object counts read "
            f"from the account on {f['verified_on']}.")
    return s


def s04_app(prs, f, shots):
    s = content(prs, f"{f['app_page_count']} pages, and nothing to set up",
                "Installs from the internal Marketplace with its data bundled in")
    picture(s, shots["overview"], LEFT, TOP, 5.9, 3.5)
    pages = f["app_pages"]
    half = (len(pages) + 1) // 2
    caption(s, LEFT + 6.1, TOP, 3.0, "every page")
    stack(s, LEFT + 6.1, TOP + 0.30, 1.5, 3.0,
          [(f"· {p}", 9.5, False, DK1, 6) for p in pages[:half]])
    stack(s, LEFT + 7.65, TOP + 0.30, 1.5, 3.0,
          [(f"· {p}", 9.5, False, DK1, 6) for p in pages[half:]])
    note(s, "Department and Company filters apply across all pages. Shown: Workforce Overview.")
    return s


def s05_workforce(prs, f, shots):
    k = f["kpis"]
    s = content(prs, "The workforce, as SAP defines it",
                "Headcount, attrition and pay on SAP's own metric definitions")
    picture(s, shots["attrition"], LEFT, TOP, 5.9, 3.5)
    x = LEFT + 6.1
    stat(s, x, TOP, 3.0, 0.82, f"{k['active_headcount']:,}", "active headcount",
         f"of {f['employees']:,} records", accent=SF_BLUE)
    stat(s, x, TOP + 0.92, 3.0, 0.82, f"{k['attrition_pct']}%", "attrition shown",
         f"{f['terminations']} terminations", accent=RED)
    stat(s, x, TOP + 1.84, 3.0, 0.82, f"${k['avg_salary']:,}", "average base salary",
         f"compa-ratio {f['avg_compa_ratio']}", accent=TEAL)
    stat(s, x, TOP + 2.76, 3.0, 0.74, f"{f['departments']}", "departments",
         f"across {f['companies']} companies", accent=DK2)
    note(s, "Attrition here is terminations as a share of all records — it is not an "
            "annualised rate, and saying so protects the number.")
    return s


def s06_equity(prs, f, shots):
    s = content(prs, "Pay equity becomes measurable",
                "Compa-ratio against range midpoint, by grade and gender")
    picture(s, shots["compensation"], LEFT, TOP, 5.9, 3.5)
    k = f["kpis"]
    card(s, LEFT + 6.1, TOP, 3.0, 1.55, "why this page matters", [
        ("Compa-ratio compares actual pay to the midpoint of the band SAP already "
         "holds — so the comparison is SAP's, not one an analyst invented.",
         10, False, DK1, 0)], accent=TEAL)
    stat(s, LEFT + 6.1, TOP + 1.68, 3.0, 0.80, f"{f['avg_compa_ratio']}",
         "average compa-ratio", "1.0 = paid at midpoint", accent=TEAL)
    stat(s, LEFT + 6.1, TOP + 2.60, 3.0, 0.80, f"{k['pct_female']}%",
         "female", f"{k['pct_managers']}% are managers", accent=VIOLET)
    note(s, "Governance belongs in this conversation: masking and row access policies "
            "decide who may see compensation at all.")
    return s


def s07_agent(prs, f, shots):
    sv = next(iter(f["semantic_view_detail"].values()))
    s = content(prs, "Plain-English questions, governed SQL underneath",
                "Eight questions ship — every one a workforce question")
    picture(s, shots["analyst"], LEFT, TOP, 5.5, 3.5)
    qs = f["agent_questions"]
    stack(s, LEFT + 5.75, TOP - 0.02, 3.4, 2.4,
          [(f"· {q}", 9.5, False, DK1, 5) for q in qs])
    card(s, LEFT + 5.75, TOP + 2.35, 3.4, 1.15, "the limit to know", [
        (f"The semantic view carries one table ({', '.join(sv['table_names'])}). "
         f"Performance, learning and recruiting are dashboard-only.",
         9.5, False, DK1, 0)], accent=RED)
    note(s, "The generated SQL expands on screen — constrained to the model, which is "
            "the reason for having one.")
    return s


def s08_regions(prs, f):
    s = content(prs, "One definition, three governed installs",
                "Region-scoped organization listing, no cross-region querying")
    y = TOP
    for r in f["regions"]:
        region = region_label(r.get("region", "—"))
        listings = ", ".join(f"{x['name']} ({x['state']})" for x in r.get("listings", [])) \
            or r.get("error", "none found")
        box(s, LEFT, y, FULLW, 0.72, LIGHT_BG, SF_BLUE)
        stack(s, LEFT + 0.28, y + 0.12, FULLW - 0.5, 0.50, [
            (region, 12.5, True, DK1, 2), (listings, 9.5, False, BODY_GREY, 0)])
        y += 0.80
    banner(s, y + 0.12, [
        ("Each region is its own governed install. For workforce data that "
         "separation is a feature — name it rather than apologise for it.",
         11.5, True, WHITE, 0)], h=0.56)
    note(s, "People 360 ships the Native App listing only — there is no separate "
            "data-share listing today, unlike Supply Chain 360.")
    return s


def s09_caveats(prs, f):
    ids = ", ".join(r["EMPLOYEE_ID"] for r in f["sample_employee_ids"])
    s = content(prs, "What is real, and what is a demo dataset",
                "Say this before you are asked — it is workforce data")
    rows = [
        ("Medallion architecture, semantic view, agent, app", "Real — this is the deliverable", TEAL, DK1),
        ("SuccessFactors field structure", "Real structure, modelled on BDC workforce products", SF_BLUE, DK1),
        ("The people, pay and ratings", f"Synthetic — demo database, IDs like {ids}", RED, RED),
        ("Names anywhere in the analytics layer", "None surfaced", TEAL, DK1),
        ("Refresh", f"{len(f['dynamic_tables'])} of {len(f['analytics_objects'])} gold tables is dynamic", RED, RED),
    ]
    y = TOP
    for label, status, accent, text_colour in rows:
        box(s, LEFT, y, FULLW, 0.50, LIGHT_BG, accent)
        stack(s, LEFT + 0.28, y + 0.13, 4.6, 0.26, [(label, 10.5, True, DK1, 0)])
        stack(s, LEFT + 5.0, y + 0.13, 4.0, 0.26, [(status, 10.5, False, text_colour, 0)])
        y += 0.56
    card(s, LEFT, y + 0.04, FULLW, 0.62, "data windows differ by table", [
        (window_line(f), 9.5, False, DK1, 0)], accent=DK2)
    note(s, "Performance reviews stop a year before the rest — the window most likely "
            "to catch someone out.")
    return s


def s10_next(prs, f):
    s = content(prs, "Where to take it", "Four options, in order of effort")
    opts = [
        ("Demo it", "Install the Native App listing in your region and present. Nothing to build."),
        ("Show the architecture", "Run the SQL in your own account and demo through "
                                  "Snowflake Intelligence — no app required."),
        ("Extend the agent", "Add performance, learning and recruiting to the semantic "
                             "view so the agent covers them too. Small, scoped, visible."),
        ("Point it at customer data", "Swap L0 for the customer's own BDC workforce "
                                      "products; the layers above still apply."),
    ]
    y = TOP
    for i, (head, detail) in enumerate(opts, 1):
        box(s, LEFT, y, FULLW, 0.68, LIGHT_BG, SF_BLUE)
        stack(s, LEFT + 0.30, y + 0.11, FULLW - 0.6, 0.46, [
            (f"{i}.  {head}", 12, True, DK2, 2), (detail, 10, False, DK1, 0)])
        y += 0.78
    banner(s, y + 0.10, [
        (f"Kit on SAP Partnership Compass · start with 03_SE_Quick_Start.docx · {f['repo']}",
         10.5, True, WHITE, 0)], h=0.50)
    return s


def main():
    f, shots = load_facts()
    prs = new_presentation()

    slides = [
        s01_cover(prs, f),
        s02_problem(prs, f),
        s03_pattern(prs, f),
        s04_app(prs, f, shots),
        s05_workforce(prs, f, shots),
        s06_equity(prs, f, shots),
        s07_agent(prs, f, shots),
        s08_regions(prs, f),
        s09_caveats(prs, f),
        s10_next(prs, f),
    ]

    issues = 0
    for n, s in enumerate(slides, 1):
        found = verify_slide(s, prs, n)
        for msg in found:
            print(f"  slide {n}: {msg}")
        issues += len(found)
    issues += len(verify_deck(prs))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"\nwrote {OUT}  ({len(prs.slides)} slides, {issues} verifier issue(s))")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
