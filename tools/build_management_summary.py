#!/usr/bin/env python3
"""Build 01_Management_Summary.docx for SAP People 360.

The customer-facing leave-behind. Unlike the SE documents it assumes no Snowflake
knowledge, and it is explicit about what is real versus synthetic — which matters
more here than in the other domains, because the subject is people.

Figures come from /tmp/people_facts.json (tools/people_facts.py).

    python3 tools/people_facts.py && python3 tools/build_management_summary.py
"""
from __future__ import annotations

import json
import pathlib
import sys
from datetime import date

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from docx_kit import (  # noqa: E402
    GREY, SAP_NAVY, SNOW_BLUE, body, bullet, callout, h1, setup_page, table,
)

KIT = pathlib.Path.home() / "Documents" / "SAP" / "People_360_Presales_Kit"
FACTS_FILE = pathlib.Path("/tmp/people_facts.json")
DATE = date.today().strftime("%d %B %Y")


def load_facts() -> dict:
    if not FACTS_FILE.exists():
        sys.exit(f"{FACTS_FILE} missing — run tools/people_facts.py first")
    return json.loads(FACTS_FILE.read_text())["facts"]


def title_block(doc, title, subtitle, strap):
    for text_, size, bold, color, after in (
        (title, 20, True, SAP_NAVY, 2),
        (subtitle, 11.5, False, SNOW_BLUE, 2),
        (strap, 9, False, GREY, 14),
    ):
        p = doc.add_paragraph()
        r = p.add_run(text_)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)


def main() -> int:
    F = load_facts()
    k = F["kpis"]
    sv = next(iter(F["semantic_view_detail"].values()))
    w = F["data_windows"]

    KIT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_page(doc)
    title_block(
        doc,
        "SAP People 360",
        "Workforce analytics on SAP Business Data Cloud and Snowflake",
        f"Management summary  ·  {DATE}",
    )

    h1(doc, "Executive summary")
    body(
        doc,
        "We have a working application that answers workforce questions in seconds "
        "which normally require an extract request and a spreadsheet: where attrition "
        "is concentrated, whether pay sits where the grade says it should, which "
        "critical roles are exposed, and how headcount really breaks down.",
    )
    body(
        doc,
        "It runs on SAP SuccessFactors data reached through SAP Business Data Cloud "
        "with zero copy. Nothing is extracted from SAP, no second copy of employee "
        "data is created, and SAP's own definitions of headcount, tenure and "
        "compa-ratio are preserved rather than rebuilt downstream. The analytics, the "
        "governed semantic layer, the natural-language agent and the application all "
        "run in Snowflake on the shared data.",
    )

    table(
        doc,
        ["Item", "Summary"],
        [
            ["Business question answered",
             "Where is workforce risk — attrition, pay inequity, succession exposure — "
             "and can we see it without moving employee data?"],
            ["Status", "Built, deployed and published in three regions. Running today."],
            ["Data foundation",
             f"SAP BDC workforce data product in Snowflake — {F['employees']:,} employee "
             f"records, {F['departments']} departments, {F['companies']} companies"],
            ["Time to an answer", "Seconds, against days for an extract-and-model cycle"],
            ["New infrastructure required", "None. Views and tables over shared data."],
            ["Employee data movement", "None. The share is read in place."],
        ],
        widths=[2.1, 4.6],
    )

    h1(doc, "What the application shows")
    body(
        doc,
        f"{F['app_page_count']} pages across the workforce lifecycle, all filterable by "
        f"department and company.",
    )
    table(
        doc,
        ["Area", "What it answers"],
        [
            ["Workforce overview", f"{k['active_headcount']:,} active headcount, "
                                   f"{k['attrition_pct']}% attrition, ${k['avg_salary']:,} average base"],
            ["Attrition", f"{F['terminations']} terminations by year, reason and department"],
            ["Compensation", f"Compa-ratio against range midpoint — average {F['avg_compa_ratio']}"],
            ["Diversity", f"{k['pct_female']}% female, mix by generation and department"],
            ["Performance", f"{sum(r['REVIEWS'] for r in F['performance_distribution']):,} reviews by rating"],
            ["Learning", f"{F['learning_completions']:,} completions"],
            ["Recruiting", ", ".join(f"{r['REQS']} {r['STATUS'].lower()}" for r in F["recruiting_by_status"])],
            ["Org and span", f"Span of control, with {F['critical_positions']} critical positions flagged"],
        ],
        widths=[1.6, 5.1],
    )

    h1(doc, "Natural-language questions, and their current limit")
    body(
        doc,
        "A Cortex Agent answers plain-English questions over the workforce layer and "
        "shows the SQL it generated, so an answer can be checked rather than trusted. "
        "Eight questions ship with the application, covering headcount, attrition, pay, "
        "diversity, tenure and critical positions.",
    )
    callout(
        doc,
        "Stated plainly",
        f"The semantic model behind the agent currently covers the workforce table only "
        f"({', '.join(sv['table_names'])}, {sv['dimensions']} dimensions and "
        f"{sv['facts']} facts). Performance, learning and recruiting are available on "
        f"the dashboards but not yet to the agent. Extending the model to cover them is "
        f"a small, well-scoped piece of work.",
    )

    h1(doc, "What is real, and what is a demonstration dataset")
    body(
        doc,
        "This matters more here than in other domains, so it is stated before it is "
        "asked.",
    )
    ids = ", ".join(r["EMPLOYEE_ID"] for r in F["sample_employee_ids"])
    table(
        doc,
        ["Element", "Status"],
        [
            ["The architecture, semantic layer and agent", "Real — this is what would be deployed"],
            ["SuccessFactors field structure", "Real — modelled on SAP BDC workforce data products"],
            ["The employees, their pay and their ratings",
             f"Synthetic. A demonstration database, with sequential identifiers ({ids}…) "
             f"rather than SAP personnel numbers"],
            ["Names in the analytics layer", "None. No employee names are surfaced anywhere"],
            ["Live connection to a production HR system", "None"],
            ["Governance design for production",
             "Not included. Masking and row access policies would be designed with you"],
        ],
        widths=[2.6, 4.1],
    )

    h1(doc, "Data currency")
    body(
        doc,
        "The dataset does not have one cut-off date; each area has its own, and the "
        "difference is worth knowing before reading a trend.",
    )
    table(
        doc,
        ["Area", "Covers"],
        [
            ["Hires", f"{w['DT_WORKFORCE_360']['min']} to {w['DT_WORKFORCE_360']['max']}"],
            ["Terminations", f"{w['DT_WORKFORCE_360.TERMINATION_DATE']['min']} to "
                             f"{w['DT_WORKFORCE_360.TERMINATION_DATE']['max']}"],
            ["Learning", f"{w['DT_LEARNING']['min']} to {w['DT_LEARNING']['max']}"],
            ["Recruiting", f"{w['DT_RECRUITING']['min']} to {w['DT_RECRUITING']['max']}"],
            ["Performance reviews", f"review years {w['DT_PERFORMANCE']['min']} to "
                                    f"{w['DT_PERFORMANCE']['max']} — a year behind the rest"],
        ],
        widths=[1.8, 4.9],
    )

    h1(doc, "Who uses this, and for what")
    table(
        doc,
        ["Role", "Question they bring", "Where they start"],
        [
            ["CHRO / HR Director", "Where is our retention risk, and what is it costing?", "Workforce Overview"],
            ["People Analytics lead", "Can I answer new questions without an extract?", "Ask the Agent"],
            ["Total Rewards", "Is pay where the grade says it should be?", "Compensation"],
            ["Talent Acquisition", "Which critical roles are exposed, and what is open?", "Recruiting, Org and Span"],
            ["Finance business partner", "What does workforce cost look like by unit?", "Headcount, Compensation"],
            ["Enterprise architect", "How does this work without copying HR data?", "BDC Sources and Lineage"],
        ],
        widths=[1.6, 3.5, 1.6],
    )

    h1(doc, "Governance, which is the real conversation")
    body(
        doc,
        "Workforce analytics fails on permissions more often than on analytics. Because "
        "the data stays in Snowflake rather than being copied into spreadsheets, access "
        "is enforced by the platform: masking policies can hide compensation from roles "
        "that should not see it, and row access policies can limit an HR business "
        "partner to their own population. That is a design exercise rather than a "
        "product feature to switch on, and it is the right next conversation.",
    )

    h1(doc, "Options from here")
    bullet(doc, "Demonstrate. It is deployed in three regions and needs no further work to show.")
    bullet(doc, "Extend the agent to performance, learning and recruiting so natural-language questions cover the full lifecycle.")
    bullet(doc, "Point it at your own workforce data products and see the same pages against your organisation.")
    bullet(doc, "Design the governance model — masking, row access, and who approves what — as the step that makes it production-ready.")

    h1(doc, "Access")
    body(
        doc,
        f"Published as an internal organization listing in three regions. A "
        f"credential-free browser build is available for review at {F['public_url']} — "
        f"it carries no live connection and no customer data. Source and documentation: "
        f"{F['repo']}.",
    )
    body(doc, f"All figures in this document were read from the platform on {F['verified_on']}.", italic=True)

    out = KIT / "01_Management_Summary.docx"
    doc.save(out)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
